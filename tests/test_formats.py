import io

import docx
import pytest

from backend.rag.ingest import IngestError, parse_upload


def minimal_pdf(text: str) -> bytes:
    """A tiny valid one-page PDF containing `text` (no external tools needed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)
    return out


def test_pdf_text_is_extracted():
    entries = parse_upload("Sports-Day_rules.pdf", minimal_pdf("Sports day is on the first Friday of November"), "public")
    assert entries[0]["title"] == "Sports Day Rules"
    assert "first Friday of November" in entries[0]["content"]


def test_pdf_without_text_gets_a_clear_error():
    with pytest.raises(IngestError, match="no readable text"):
        parse_upload("scan.pdf", minimal_pdf(""), "public")
    with pytest.raises(IngestError):
        parse_upload("broken.pdf", b"not a pdf", "public")


def test_docx_paragraphs_and_tables():
    d = docx.Document()
    d.add_paragraph("Uniform policy: white shirt and navy shorts.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Shoes", "Black"
    buf = io.BytesIO()
    d.save(buf)
    text = parse_upload("uniform.docx", buf.getvalue(), "student")[0]["content"]
    assert "white shirt" in text and "Shoes | Black" in text


def test_text_markdown_and_errors():
    assert parse_upload("notes.md", "# Title\nBody".encode(), "staff")[0]["access"] == "staff"
    with pytest.raises(IngestError):
        parse_upload("empty.txt", b"   ", "public")
    with pytest.raises(IngestError):
        parse_upload("data.json", b'[{"title": "x"}]', "public")  # missing content
    with pytest.raises(IngestError):
        parse_upload("data.json", b'[{"title":"x","content":"y","access":"everyone"}]', "public")
    with pytest.raises(IngestError):
        parse_upload("file.exe", b"MZ", "public")

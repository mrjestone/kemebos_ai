"""Turn an uploaded file or typed entry into searchable chunks."""
import io
import json
from datetime import datetime, timezone

from bson import ObjectId
from pymongo.database import Database

from .chunking import chunk_text
from .embeddings import embed_documents

MAX_CHUNKS = 300      # protects the free embedding quota from a runaway file
MAX_CHARS = 400_000


class IngestError(ValueError):
    """A problem the admin can fix (bad file, empty text, too big)."""


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        text = "\n\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    except Exception as exc:
        raise IngestError("Couldn't read this PDF.") from exc
    if len(text.strip()) < 20:
        raise IngestError("This PDF has no readable text (it may be a scan). Upload a text version instead.")
    return text


def _docx(data: bytes) -> str:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestError("Couldn't read this Word file.") from exc
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n\n".join(parts)


def _title_from(filename: str) -> str:
    stem = filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").strip()
    return stem.title() or "Untitled"


def parse_upload(filename: str, data: bytes, default_access: str) -> list[dict]:
    """Returns entries: [{title, content, access}]. JSON files may hold many."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "json":
        try:
            raw = json.loads(data.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise IngestError("Invalid JSON file.") from exc
        items = raw.get("entries") if isinstance(raw, dict) else raw
        if not isinstance(items, list) or not items:
            raise IngestError('JSON must be a list of {"title","content","access"} entries.')
        entries = []
        for i, it in enumerate(items, 1):
            if not isinstance(it, dict) or not str(it.get("title", "")).strip() or not str(it.get("content", "")).strip():
                raise IngestError(f"Entry {i} needs a title and content.")
            access = it.get("access", default_access)
            if access not in ("public", "student", "staff", "admin"):
                raise IngestError(f"Entry {i} has an invalid access level.")
            entries.append({"title": str(it["title"]).strip(), "content": str(it["content"]).strip(), "access": access})
        return entries
    if ext in ("txt", "md"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise IngestError("Text files must be UTF-8.") from exc
    elif ext == "pdf":
        text = _pdf(data)
    elif ext == "docx":
        text = _docx(data)
    else:
        raise IngestError("Use JSON, MD, TXT, PDF or DOCX.")
    if not text.strip():
        raise IngestError("The file is empty.")
    return [{"title": _title_from(filename), "content": text.strip(), "access": default_access}]


def _index_chunks(db: Database, doc_id, title: str, content: str, access: str) -> int:
    if len(content) > MAX_CHARS:
        raise IngestError("This document is too long. Split it into smaller files.")
    chunks = chunk_text(content)
    if not chunks:
        raise IngestError("Nothing to index.")
    if len(chunks) > MAX_CHUNKS:
        raise IngestError("This document is too long. Split it into smaller files.")
    # The title is embedded with the text so "Fee Payment" chunks match fee questions.
    vectors = embed_documents([f"{title}\n{c}" for c in chunks])
    db.chunks.delete_many({"docId": doc_id})
    db.chunks.insert_many([
        {"docId": doc_id, "index": i, "title": title, "text": c, "access": access, "embedding": v}
        for i, (c, v) in enumerate(zip(chunks, vectors))
    ])
    return len(chunks)


def create_document(db: Database, *, title: str, content: str, access: str, kind: str,
                    filename: str | None = None, size: int | None = None) -> dict:
    doc = {
        "title": title, "content": content, "access": access, "kind": kind,
        "filename": filename, "size": size if size is not None else len(content.encode()),
        "status": "processing", "chunks": 0, "error": None,
        "createdAt": datetime.now(timezone.utc),
    }
    doc_id = db.documents.insert_one(doc).inserted_id
    try:
        n = _index_chunks(db, doc_id, title, content, access)
    except Exception as exc:  # leave a visible, retryable record instead of half-indexed data
        db.chunks.delete_many({"docId": doc_id})
        db.documents.update_one({"_id": doc_id}, {"$set": {"status": "failed", "error": str(exc)[:200]}})
        raise
    db.documents.update_one({"_id": doc_id}, {"$set": {"status": "ready", "chunks": n}})
    return db.documents.find_one({"_id": doc_id})


def reindex_document(db: Database, doc: dict) -> dict:
    """Re-chunk and re-embed (after editing the text/title, or to retry a failed one)."""
    n = _index_chunks(db, doc["_id"], doc["title"], doc["content"], doc["access"])
    db.documents.update_one({"_id": doc["_id"]}, {"$set": {"status": "ready", "chunks": n, "error": None}})
    return db.documents.find_one({"_id": doc["_id"]})


def set_access(db: Database, doc_id: ObjectId, access: str) -> None:
    """Changing who can see a document needs no re-embedding, only the tag."""
    db.documents.update_one({"_id": doc_id}, {"$set": {"access": access}})
    db.chunks.update_many({"docId": doc_id}, {"$set": {"access": access}})


def delete_document(db: Database, doc_id: ObjectId) -> None:
    db.chunks.delete_many({"docId": doc_id})
    db.documents.delete_one({"_id": doc_id})
